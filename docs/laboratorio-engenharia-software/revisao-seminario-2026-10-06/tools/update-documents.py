from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
R=Path(__file__).resolve().parents[1]; OLD=R.parent/'revisao-astah-2026-09-28'
D=json.loads((R/'tools/revision-plan.json').read_text(encoding='utf8'))
OUT=R/'documentos';OUT.mkdir(exist_ok=True)
def pic(id):
 matches=[p for p in (R/'diagramas').rglob('*.png') if p.stem.startswith(id+' ')]
 assert len(matches)==1,(id,matches)
 return matches[0]
def para(d,t='',style=None,before=None):
 p=d.add_paragraph(t,style)
 if before is not None:before.addprevious(p._p)
 return p
def heading(d,t,l=2,before=None,page=False):
 p=para(d,t,'Heading '+str(l),before);p.paragraph_format.page_break_before=page;return p
def image(d,path,before=None,maxw=6.4,maxh=7.3):
 p=para(d,before=before);p.alignment=1
 w,h=Image.open(path).size;s=min(maxw/w,maxh/h)
 p.add_run().add_picture(str(path),width=Inches(w*s),height=Inches(h*s));p.paragraph_format.keep_with_next=True;return p
def figure(d,id,before=None,maxh=7.3):
 image(d,pic(id),before,maxh=maxh)
 p=para(d,'Fonte editável: Stacklyst-seminario.asta / '+id+'.',before=before)
 p.alignment=1
 for r in p.runs:r.font.size=Pt(9)
def allparas(d):
 return list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
def replace(p,new):
 if p.text==new:return
 # Preserve paragraph/style and first-run appearance for local text edits.
 if p.runs:p.runs[0].text=new;[setattr(r,'text','') for r in p.runs[1:]]
 else:p.add_run(new)
def apply_style(d):
 for name,size in [('Normal',11),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
  s=d.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 for p in d.paragraphs:
  if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
 for t in d.tables:
  for row in t.rows:
   for cell in row.cells:
    for p in cell.paragraphs:
     for run in p.runs:run.font.size=Pt(9)
def add_table(d,headers,rows,before=None):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for values in rows:
  for c,v in zip(t.add_row().cells,values):c.text=str(v)
 if before is not None:before.addprevious(t._tbl)
 return t

# Document 02: retain the supplied cover and document styles, replace obsolete process content.
d=Document(OLD/'entrega-final/02-atividades-do-negocio-stacklyst.docx')
start=next(p._p for p in d.paragraphs if p.text.startswith('2. Objetivo'))
body=d._element.body;children=list(body);i=children.index(start)
for e in children[i:]:
 if e.tag!=qn('w:sectPr'):body.remove(e)
for p in d.paragraphs:replace(p,p.text.replace('28/09/2026','06/10/2026'))
heading(d,'2. Objetivo',1)
para(d,'Este documento descreve como estudantes, candidatos, representantes de empresas e recrutadores realizam suas atividades antes de usar o Stacklyst. O foco é o processo atual: estudar com materiais dispersos, conferir exercícios, divulgar oportunidades e acompanhar contatos. Os diagramas identificam os responsáveis em raias e mostram decisões e retornos do processo.')
para(d,'A revisão segue as orientações de 28/09/2026 consolidadas no documento da Profª Patrícia e o capítulo 10 do livro UML 2 Uma abordagem prática, de Gilleanes T. A. Guedes. Os exemplos descrevem processos possíveis e coerentes com a orientação, sem pretender retratar uma pesquisa de campo que não foi realizada.')
heading(d,'3. Atividades do negócio',1)
para(d,'AN01, AN03, AN04 e AN06 tratam do aprendizado. AN08A e AN08B tratam da divulgação e do acompanhamento de oportunidades. AN09 trata da descoberta e participação em eventos. Os identificadores anteriores de atividades exclusivas do sistema foram retirados desta versão, sem renumerar os processos mantidos.')
for i,a in enumerate(D['activities'],1):
 heading(d,f'3.{i} {a["id"]} {a["name"]}',2,page=True)
 para(d,'Objetivo: '+a['purpose']);para(d,'Participantes: '+', '.join(a['lanes'])+'.')
 para(d,'Situação atual: '+a['difficulty'])
 heading(d,'Fluxo e decisões',3)
 for z in a['nodes']:
  if z[1]=='action':para(d,a['lanes'][z[7]]+': '+z[2].replace('\n',' ')+'.')
 nodes={z[0]:z for z in a['nodes']}
 for e in a['edges']:
  if len(e)>2 and e[2]:para(d,'Decisão '+nodes[e[0]][2].replace('\n',' ')+f' [{e[2]}]: '+(nodes[e[1]][2].replace('\n',' ') or 'encerrar o processo')+'.')
 para(d,'Resultado: '+a['result'])
 heading(d,a['id']+' Diagrama do processo atual',3,page=True);figure(d,a['id'],maxh=6.7)
heading(d,'4. Relação com o sistema proposto',1,page=True)
para(d,'Os processos de negócio contextualizam os problemas que o sistema procura apoiar. Os casos de uso do documento 04 descrevem a interação com o Stacklyst. O duelo continua como funcionalidade do sistema, enquanto o processo de estudo atual utiliza ações humanas como pesquisar, praticar e revisar.')
add_table(d,['Processo atual','Apoio do Stacklyst'],[
 ['AN01 Estudar programação','UC004 Selecionar trilha; UC005 Estudar lição'],['AN03 Praticar exercícios','UC006 Resolver exercício; UC028 Avaliar resposta; UC029 Executar código e testes'],['AN04 Acompanhar aprendizagem','UC007 Consultar progresso'],['AN06 Buscar ajuda','UC012 Consultar ajuda do exercício'],['AN08A Divulgar oportunidade','UC017 Publicar vaga; UC025 Cadastrar empresa'],['AN08B Acompanhar seleção','UC018 Consultar vagas; UC022 Candidatar-se; UC026 Atualizar etapa; UC056 Consultar candidaturas'],['AN09 Participar de evento','UC046 Consultar eventos; UC023 Inscrever-se ou cancelar participação']])
heading(d,'5. Referências',1)
para(d,'Anotações das Orientações da Profª Patrícia. Projeto Stacklyst. 28/09/2026. Documento de apoio fornecido pela equipe, que consolida orientações verbais.')
para(d,'GUEDES, Gilleanes T. A. UML 2 Uma abordagem prática. Arquivo PDF fornecido. Seções 4.1, 10.6 e 10.22. Partições identificam responsabilidades; decisões representam alternativas e junções sincronizam atividades paralelas.')
apply_style(d);d.save(OUT/'02-atividades-do-negocio-stacklyst.docx')

# Document 04: retain specifications and UI evidence, correct IDs, scope and diagram sources.
d=Document(OLD/'fontes-drive/04-casos-de-uso-stacklyst-atualizado.docx')
by_name={u['name']:u['id'] for u in D['uc']};by_name['Participar de evento']='UC023'
for p in d.paragraphs:
 t=p.text
 if 'Caso de Uso UC' in t:
  title=t.split(' — ',1)[-1];uid=by_name.get(title)
  if uid:replace(p,f'5.{int(uid[2:])} Caso de Uso {uid} — '+next(u['name'] for u in D['uc'] if u['id']==uid))
 for old,new in [('28/09/2026','06/10/2026'),('Visitante','Usuário'),('visitante','usuário'),('Mensagens e comunidade','Mensagens privadas')]:
  t=p.text;replace(p,t.replace(old,new))
 if t.startswith('Este documento especifica'):replace(p,'Este documento especifica 52 casos de uso do escopo atual do Stacklyst, com atores, condições, fluxos e referências de implementação. Os identificadores UC001 a UC056 permanecem iguais aos do Astah, exceto UC042 a UC045, retirados com as comunidades. A visão geral agrupa a cobertura em 16 capacidades. Os diagramas são nativos e editáveis no Astah UML.')
 if t.startswith('O panorama UC00 reúne'):replace(p,'O panorama UC00 agrupa os 52 casos de uso em 16 capacidades G01 a G16. Os 11 recortes seguintes mostram os casos detalhados com os mesmos identificadores. A tabela de cobertura relaciona cada agrupamento a seus casos. Agrupamento visual não representa uma relação UML include.')
 if t.startswith('Panorama geral resumido'):replace(p,'Visão geral UC00 com 16 agrupamentos e cobertura dos 52 casos de uso mantidos.')
 if t.startswith('Limite de UC045:'):replace(p,'Comunidades estão fora do escopo entregue. O código e estruturas legadas podem continuar no repositório, sem compor os casos de uso apresentados.')
 if t.startswith('A especificação textual e os diagramas abrangem'):replace(p,'A especificação e os diagramas usam os mesmos 52 casos. A revisão de 06/10/2026 confere o código do commit 9256a4f e os modelos salvos do Astah. Esta rastreabilidade documenta a implementação e não substitui testes de aceitação.')
 if t.startswith('12 conceitos centrais;'):replace(p,'Classes e módulos responsáveis pelas operações principais. Os métodos exibidos correspondem a funções reais do código.')
 if t.startswith('As 10 vistas correspondentes'):replace(p,'As vistas CL01 a CL12 do Astah mostram entidades e os serviços ou módulos responsáveis por suas operações. As entidades Prisma conservam atributos e associações. Os métodos pertencem aos objetos de serviço e módulos TypeScript identificados na matriz de operações.')
 if 'Fonte editável:' in t:replace(p,p.text.replace('Stacklyst-revisado-simplificado.asta','Stacklyst-seminario.asta'))
# Correct protected preconditions without applying authentication to public reads or internal include cases.
current=None;public={'UC001','UC002','UC011','UC018','UC028','UC029','UC031','UC032','UC046','UC050'}
for p in d.paragraphs:
 mt=re.search(r'Caso de Uso (UC\d+)',p.text)
 if mt:current=mt[1]
 if p.text.startswith('Pré-condições') and current not in public and 'Usuário deve estar logado' not in p.text:
  label,_,rest=p.text.partition(':');replace(p,label+': Usuário deve estar logado. '+rest.strip())
# Align UC023 with both operations in the route.
current=None
for p in d.paragraphs:
 mt=re.search(r'Caso de Uso (UC\d+)',p.text)
 if mt:current=mt[1]
 if current=='UC023' and p.text.startswith('Descrição:'):replace(p,'Descrição: Inscrever o usuário em um evento ou cancelar sua própria participação, conforme as condições do evento.')
 if current=='UC023' and p.text.startswith('Pós-condições:'):replace(p,'Pós-condições: Participação registrada ou removida. O cancelamento não é permitido para evento concluído.')
# Replace all Astah figures using the adjacent source reference, preserving UI screenshots.
paras=list(d.paragraphs)
replaced=[]
for i,p in enumerate(paras):
 mt=re.search(r'Fonte editável:.* / ((?:UC|CL|MR|TC)\d{2})\.',p.text)
 if not mt:continue
 ident=mt[1];imagep=next((q for q in reversed(paras[max(0,i-4):i]) if q._p.xpath('.//w:drawing')),None)
 if imagep is not None:
  for e in list(imagep._p):
   if e.tag!=qn('w:pPr'):imagep._p.remove(e)
  path=pic(ident);w,h=Image.open(path).size;s=min(6.35/w,7.05/h)
  imagep.add_run().add_picture(str(path),width=Inches(w*s),height=Inches(h*s));replaced.append(ident)
# Clean obsolete community records, preserve non-community collaboration and feed.
for t in list(d.tables):
 for row in list(t.rows)[1:]:
  joined=' '.join(c.text for c in row.cells)
  if re.search(r'Consultar comunidades|Criar comunidade|Entrar ou sair de comunidade|Administrar comunidade|Comunidades',joined):t._tbl.remove(row._tr)
for p in allparas(d):
 t=p.text
 if 'todas as 56 funções' in t:replace(p,'Critério adotado: 52 casos de uso mantêm os IDs do Astah. UC042 a UC045 foram retirados com as comunidades. As 16 capacidades da visão geral cobrem todos os casos detalhados.')
 elif 'comunidade' in t.lower() and t.startswith('Melhoria futura:'):replace(p,t.replace('comunidade, ','publicações, '))
 # Correct complementary table IDs by the functionality name, not a blind numeric substitution.
for t in d.tables:
 for row in t.rows:
  if len(row.cells)>1:
   uid=by_name.get(row.cells[1].text.strip())
   if uid and row.cells[0].text.strip().startswith('UC'):row.cells[0].text=uid
# Replace obsolete business-to-system matrix with current processes.
matrix=next(t for t in d.tables if t.cell(0,0).text=='Atividade do Negócio')
for row in list(matrix.rows)[1:]:matrix._tbl.remove(row._tr)
mapping=[('AN01 Estudar programação','RF005–RF010','RN001–RN002','UC004, UC005'),('AN03 Praticar exercícios','RF005–RF010','RN025, RN027–RN029','UC006, UC028, UC029'),('AN04 Acompanhar aprendizagem','RF005–RF010','RN001–RN002','UC007'),('AN06 Buscar ajuda','RF005–RF010','RN025','UC012'),('AN08A Divulgar oportunidade','RF035, RF038','Conforme regras de vagas','UC017, UC025'),('AN08B Acompanhar seleção','RF035, RF038, RF048','Conforme regras de candidaturas','UC018, UC022, UC026, UC056'),('AN09 Descobrir evento','RF041','Conforme regras de eventos','UC023, UC046')]
for vals in mapping:
 for c,v in zip(matrix.add_row().cells,vals):c.text=v
# Add explicit source-backed scope decisions beside the existing consistency discussion.
anchor=next(p._p for p in d.paragraphs if p.text.startswith('9. Modelos'))
heading(d,'8.1 Escopo apresentado no seminário',2,anchor)
para(d,'UC057 Notificar usuário foi retirado como objetivo iniciado manualmente por ator. A notificação pertence ao fluxo que a dispara e a NotificationService.create. Em atualização de candidatura, o destinatário é application.user_id. UC058 Manter intenção de vaga e UC059 Manter interesse em evento foram retirados por não haver um cadastro separado correspondente no código conferido. Permanecem UC022 Candidatar-se a vaga e UC023 Inscrever-se ou cancelar participação em evento.',before=anchor)
para(d,'UC011 mantém o nome Consultar ranking de XP: a ordenação e a concessão de pontos são automáticas. UC013 e UC035 detalham criação e edição/exclusão de publicação própria, reunidas em G06 na visão geral. O feed e as mensagens privadas permanecem no escopo.',before=anchor)
heading(d,'8.2 Cobertura da visão geral',2,anchor)
add_table(d,['Agrupamento','Capacidade','Casos detalhados'],[(g[0],g[1].replace('\n',' '),', '.join(f'UC{x:03}' for x in g[2])) for g in D['overview']],anchor)
heading(d,'8.3 Operações e responsáveis',2,anchor,page=True)
para(d,'Os objetos de serviço e módulos abaixo são representados como classificadores UML com operações. A notação não afirma que todos sejam declarações class em TypeScript. Os nomes dos métodos são literais; as assinaturas gráficas são resumidas. Operações não são acrescentadas artificialmente aos registros Prisma.',before=anchor)
add_table(d,['Responsável','Operações reais','Arquivo'],[(s['name'],'\n'.join(o[0]+'()' for o in s['ops']),s['path']) for s in D['services']],anchor)
for ident,title in [('CL08','Classes de vagas e candidaturas'),('CL09','Classes de eventos'),('CL03','Classes de prática e progresso')]:
 heading(d,title,2,page=True);figure(d,ident)
apply_style(d);d.save(OUT/'04-casos-de-uso-stacklyst.docx')
print('Updated documents, replaced',len(replaced),'Astah figures')

# The existing vision/requirements describe the feed as "community". Use an
# unambiguous label and retain the original requirements, logos and tables.
for prefix in ['01','03']:
 path=next((OLD/'entrega-final').glob(prefix+'*.docx'));d=Document(path)
 for p in allparas(d):
  t=p.text
  t=t.replace('duelos, comunidade e acompanhamento','duelos, publicações, mensagens e acompanhamento')
  t=t.replace('aprendizagem, comunidade e oportunidades','aprendizagem, interação entre pessoas e oportunidades')
  t=t.replace('Grupo 7 — Comunidade','Grupo 7 — Publicações e interação')
  t=t.replace('denúncias da comunidade','denúncias de publicações')
  if t.strip()=='Comunidade':t='Publicações e interação'
  replace(p,t)
 anchor=next((p._p for p in d.paragraphs if p.style.name=='Heading 1' and not p.text.startswith('1.')),None)
 if anchor is not None:
  para(d,'Escopo documental de 06/10/2026: o feed, as publicações e as mensagens privadas permanecem no sistema apresentado. A funcionalidade de comunidades foi retirada do escopo. As atividades de negócio e os casos de uso correspondem à revisão do Astah e aos documentos 02 e 04 desta entrega.',before=anchor)
 d.save(OUT/path.name)
