import json, re, zipfile
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
REPO=ROOT.parents[2]
data={}
for f in (ROOT/'fontes-drive').glob('*.docx'):
    d=Document(f)
    lines=[p.text for p in d.paragraphs]
    tables=[[ [c.text for c in row.cells] for row in t.rows] for t in d.tables]
    data[f.name[:2]]={'paragraphs':lines,'tables':tables}
    f.with_suffix('.txt').write_text('\n'.join(lines+[' | '.join(r) for t in tables for r in t]),encoding='utf8')
data['uc']=[{'id':m[1],'name':m[2]} for p in data['04']['paragraphs'] if (m:=re.match(r'5\.\d+ Caso de Uso (UC\d+) — (.+)',p))]
schema=(REPO/'prisma/schema.prisma').read_text(encoding='utf8')
models={n:b for n,b in re.findall(r'model (\w+) \{(.*?)\n\}',schema,re.S)}
enums={n:re.findall(r'^\s*(\w+)\s*$',b,re.M) for n,b in re.findall(r'enum (\w+) \{(.*?)\n\}',schema,re.S)}
data['models']=[]; data['relations']=[]; data['enums']=enums
for name,body in models.items():
    fields=[]
    for line in body.splitlines():
        m=re.match(r'\s*(\w+)\s+([\w?\[\]]+)(.*)',line)
        if not m:continue
        fn,typ,rest=m.groups(); base=typ.rstrip('?[]')
        if base in models:
            rel=re.search(r'fields:\s*\[([^]]+)\].*?references:\s*\[([^]]+)\]',rest)
            if rel:
                keys=[x.strip() for x in rel[1].split(',')]
                unique=any(re.search(r'^\s*'+re.escape(k)+r'\s+[^\n]*@unique',body,re.M) for k in keys) or bool(re.search(r'@@(?:unique|id)\(\[\s*'+r'\s*,\s*'.join(keys)+r'\s*\]',body))
                data['relations'].append({'from':name,'to':base,'role':fn,'fk':keys,'ref':[x.strip() for x in rel[2].split(',')],'targetMultiplicity':'0..1' if '?' in typ else '1','sourceMultiplicity':'0..1' if unique else '0..*','definition':rest.strip()})
        else:fields.append({'name':fn,'type':typ,'pk':'@id' in rest,'unique':'@unique' in rest,'definition':rest.strip()})
    constraints=[l.strip() for l in body.splitlines() if l.strip().startswith('@@')]
    primary=re.search(r'@@id\(\[([^]]+)\]',body)
    if primary:
        keys=[x.strip() for x in primary[1].split(',')]
        for f in fields:f['pk']=f['pk'] or f['name'] in keys
    if name=='ExerciseSubmission':constraints.append('UNIQUE(user_id, exercise_id) WHERE first_completion = true; migration 20260822090000')
    mapped=re.search(r'@@map\("([^"]+)"\)',body)
    data['models'].append({'name':name,'table':mapped[1] if mapped else name,'fields':fields,'constraints':constraints})
for m in data['models']:
    fks={f for r in data['relations'] if r['from']==m['name'] for f in r['fk']}
    for f in m['fields']:f['fk']=f['name'] in fks
(ROOT/'tools/source-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf8')
print(f"Extracted {len(data['uc'])} use cases, {len(models)} tables, {len(data['relations'])} foreign keys")
