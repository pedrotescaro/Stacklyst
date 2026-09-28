from pathlib import Path
from docx import Document
R=Path(__file__).resolve().parents[1]
for num in ['02','04']:
 d=Document(next((R/'fontes-drive').glob(num+'*.docx')))
 lines=[]
 for i,p in enumerate(d.paragraphs):
  if p.text or p._p.xpath('.//w:drawing'):lines.append(f'{i}: {p.style.name} | {p.text[:360]}'+(' [IMAGE]' if p._p.xpath('.//w:drawing') else ''))
 (R/'tools'/f'paragraphs-{num}.txt').write_text('\n'.join(lines),encoding='utf8')
 print(num,len(d.paragraphs),len(d.tables))
